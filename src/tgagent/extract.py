"""Text extraction from common document formats."""

from __future__ import annotations

from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".json", ".log", ".py", ".xml", ".html", ".yaml", ".yml", ".csv"}


def _truncate(text: str, limit: int) -> str:
    """Keep head and tail of an oversized extraction."""
    if len(text) <= limit:
        return text
    half = limit // 2
    return f"{text[:half]}\n\n... [truncated {len(text) - limit} chars] ...\n\n{text[-half:]}"


def _parse_pages(spec: str) -> set[int] | None:
    """Parse a page spec like '1-5,8' into a set of 1-based page numbers."""
    if not spec or not spec.strip():
        return None
    pages: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            lo, _, hi = part.partition("-")
            pages.update(range(int(lo), int(hi) + 1))
        else:
            pages.add(int(part))
    return pages


def _extract_pdf(path: Path, pages: set[int] | None) -> str:
    import pdfplumber

    chunks: list[str] = []
    texts: list[str] = []
    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        wanted = sorted(p for p in pages if 1 <= p <= total) if pages else range(1, total + 1)
        for number in wanted:
            text = pdf.pages[number - 1].extract_text() or ""
            texts.append(text)
            chunks.append(f"[page {number}]\n{text}")
    # A scanned PDF has no text on ANY page; a short-but-real PDF still extracts something.
    if not pages and texts and not any(t.strip() for t in texts):
        return (
            "PDF appears to be scanned (no text layer): extraction returned no text on any page. "
            "Tell the user OCR would be needed."
        )
    return "\n\n".join(chunks)


def _extract_docx(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path, max_rows: int = 500) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        lines = [f"[sheet: {sheet.title}]"]
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= max_rows:
                lines.append(f"... (first {max_rows} rows shown)")
                break
            lines.append(" | ".join("" if v is None else str(v) for v in row))
        chunks.append("\n".join(lines))
    workbook.close()
    return "\n".join(chunks)


def _extract_csv(path: Path, max_rows: int = 500) -> str:
    import csv

    with path.open(newline="", encoding="utf-8", errors="replace") as fh:
        rows = list(csv.reader(fh))
    out = "\n".join(",".join(row) for row in rows[:max_rows])
    if len(rows) > max_rows:
        out += f"\n... ({len(rows)} rows total, first {max_rows} shown)"
    return out


def extract_text(path: Path, pages: str = "", limit: int = 200_000) -> str:
    """Extract text from a document by suffix; best-effort, never raises for content reasons."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path, _parse_pages(pages))
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix == ".xlsx":
        text = _extract_xlsx(path)
    elif suffix in TEXT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        return f"ERROR: unsupported file type {suffix or '(no extension)'}"
    return _truncate(text, limit)
