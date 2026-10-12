"""Tests for document text extraction (fixtures generated in-test)."""

from __future__ import annotations

from pathlib import Path

from tgagent.extract import _parse_pages, _truncate, extract_text


def _make_pdf(path: Path, text: str) -> None:
    import matplotlib

    matplotlib.use("Agg")
    matplotlib.rcParams["pdf.fonttype"] = 42  # TrueType fonts keep text extractable
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=(4, 2))
    fig.text(0.1, 0.5, text)
    fig.savefig(path, format="pdf")
    plt.close(fig)


def _make_docx(path: Path, text: str) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    document.save(str(path))


def _make_xlsx(path: Path, rows: list[tuple]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_extract_txt(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("hello extraction", encoding="utf-8")
    assert "hello extraction" in extract_text(path)


def test_extract_csv(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    text = extract_text(path)
    assert "a,b" in text and "3,4" in text


def test_extract_pdf(tmp_path):
    path = tmp_path / "doc.pdf"
    _make_pdf(path, "Hello PDF Extraction 123")
    text = extract_text(path)
    assert "Hello PDF Extraction 123" in text
    assert "[page 1]" in text


def test_extract_pdf_without_text_layer(tmp_path):
    # An image-only PDF (no text elements) should trigger the scanned-PDF notice.
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    fig = plt.figure()
    plt.imshow(np.zeros((10, 10)))
    plt.axis("off")
    fig.savefig(tmp_path / "scan.pdf", format="pdf")
    plt.close(fig)

    text = extract_text(tmp_path / "scan.pdf")
    assert "scanned" in text


def test_extract_docx(tmp_path):
    path = tmp_path / "doc.docx"
    _make_docx(path, "Hello from DOCX")
    assert "Hello from DOCX" in extract_text(path)


def test_extract_xlsx(tmp_path):
    path = tmp_path / "sheet.xlsx"
    _make_xlsx(path, [("name", "value"), ("alpha", 42)])
    text = extract_text(path)
    assert "alpha | 42" in text


def test_unsupported_suffix(tmp_path):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK")
    assert extract_text(path).startswith("ERROR: unsupported file type")


def test_truncate_keeps_head_and_tail():
    text = "A" * 500 + "MIDDLE" + "B" * 500
    out = _truncate(text, 100)
    assert len(out) < len(text)
    assert out.startswith("A") and out.endswith("B")
    assert "truncated" in out


def test_parse_pages():
    assert _parse_pages("") is None
    assert _parse_pages("1-3,5") == {1, 2, 3, 5}
    assert _parse_pages("2") == {2}
